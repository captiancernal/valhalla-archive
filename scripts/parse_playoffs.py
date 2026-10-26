"""Parse the CFB Playoff Tree sheet into a list of playoff games per season.

Layout (one block per season, year in column I):
  2023: 4-team bracket. 2024+: 12-team bracket, left half in columns B-H, right half in J-P.
  Rows are relative to the year row s:
    first round  : left  B/C rows s+3 vs s+5, s+9 vs s+11 ; right P/O same rows
    quarterfinal : left  D/E rows s+4 vs s+6, s+8 vs s+10 ; right N/M same rows
    semifinal    : left  F/G rows s+6 vs s+8              ; right L/K same rows
    final        : teams H/J row s+7, scores H/J row s+6
    champion     : column I row s+3
Read-only. Usage: py -P scripts/parse_playoffs.py "<file>.xlsx" [out.json]
"""
import json
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as ci


def val(ws, col, row):
    v = ws.cell(row, ci(col)).value
    if isinstance(v, str):
        v = v.strip() or None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return v


def game(ws, rnd, tcol, scol, r1, r2):
    a, b = val(ws, tcol, r1), val(ws, tcol, r2)
    sa, sb = val(ws, scol, r1), val(ws, scol, r2)
    if not a and not b:
        return None
    g = {"round": rnd, "team_a": a, "score_a": sa, "team_b": b, "score_b": sb}
    if isinstance(sa, int) and isinstance(sb, int) and sa != sb:
        g["winner"] = a if sa > sb else b
    return g


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["CFB Playoff Tree"]
    years = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, ci("I")).value
        if isinstance(v, (int, float)) and 2000 < v < 2100:
            years.append((int(v), r))
    seasons = []
    for year, s in years:
        games = []
        if year == 2023:
            games += [g for g in [
                game(ws, "semifinal", "F", "G", s + 6, s + 8),
                game(ws, "semifinal", "L", "K", s + 6, s + 8),
            ] if g]
            fin = {"round": "final", "team_a": val(ws, "H", s + 7), "score_a": val(ws, "H", s + 6),
                   "team_b": val(ws, "J", s + 7), "score_b": val(ws, "J", s + 6)}
        else:
            for side, (fr_t, fr_s, qf_t, qf_s, sf_t, sf_s) in {
                "left": ("B", "C", "D", "E", "F", "G"),
                "right": ("P", "O", "N", "M", "L", "K"),
            }.items():
                games += [g for g in [
                    game(ws, "first_round", fr_t, fr_s, s + 3, s + 5),
                    game(ws, "first_round", fr_t, fr_s, s + 9, s + 11),
                    game(ws, "quarterfinal", qf_t, qf_s, s + 4, s + 6),
                    game(ws, "quarterfinal", qf_t, qf_s, s + 8, s + 10),
                    game(ws, "semifinal", sf_t, sf_s, s + 6, s + 8),
                ] if g]
            fin = {"round": "final", "team_a": val(ws, "H", s + 7), "score_a": val(ws, "H", s + 6),
                   "team_b": val(ws, "J", s + 7), "score_b": val(ws, "J", s + 6)}
        if fin["team_a"] or fin["team_b"]:
            sa, sb = fin["score_a"], fin["score_b"]
            if isinstance(sa, int) and isinstance(sb, int) and sa != sb:
                fin["winner"] = fin["team_a"] if sa > sb else fin["team_b"]
            games.append(fin)
        champ = val(ws, "I", s + 3)
        byes = val(ws, "Q", s + 4) if year != 2023 else None
        if games or champ:
            seasons.append({"season": year, "champion_cell": champ, "byes": byes, "games": games})
    return seasons


if __name__ == "__main__":
    data = parse(sys.argv[1])
    for s in data:
        fin = s["games"][-1] if s["games"] else {}
        print(s["season"], "champion cell:", s["champion_cell"], "| final:", fin.get("team_a"), fin.get("score_a"),
              "-", fin.get("team_b"), fin.get("score_b"), "->", fin.get("winner"), "| games:", len(s["games"]))
    if len(sys.argv) > 2:
        with open(sys.argv[2], "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)
