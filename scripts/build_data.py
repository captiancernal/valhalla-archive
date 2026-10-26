"""Build the Valhalla Archive Site Database from a tracker workbook.

Read-only on the workbook. Applies owner-approved fixes (data/fixes.json) and
canonical names (data/aliases.json), then writes one JSON file per table to
data/tables/ plus a raw cell snapshot to data/snapshot/.

Usage:
  py -P scripts/build_data.py                      # newest .xlsx in source/
  py -P scripts/build_data.py "source/<file>.xlsx"
"""
import datetime
import glob
import hashlib
import json
import os
import re
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict

import openpyxl
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_playoffs import parse as parse_playoff_tree  # noqa: E402

warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(DATA, "tables")
SNAP = os.path.join(DATA, "snapshot")

PROGRAMS = ["CLEV", "OHIO", "M-OH"]
COACH_SHEETS = {"CLEV": "Coach Josh Dolan", "OHIO": "Coach Tim Slocombe ", "M-OH": "Coach Joe Snyderwine"}
SCHOOLS = {"CLEV": "Cleveland State", "OHIO": "Ohio", "M-OH": "Miami (OH)"}
NICKNAMES = {"CLEV": "Vikings", "OHIO": "Bobcats", "M-OH": "RedHawks"}
CONF_NAMES = {"B12": "Big 12", "B10": "Big Ten", "P12": "Pac-12", "CUSA": "C-USA", "MAC": "MAC", "Sun Belt": "Sun Belt",
              "AAC": "AAC", "Mountain West": "Mountain West", "SEC": "SEC", "ACC": "ACC"}
STANDINGS_CONFS = ["AAC", "ACC", "Big 12", "Big Ten", "C-USA", "Independents", "MAC", "Mountain West", "Pac-12", "SEC", "Sun Belt"]
AWARDS_INFO = {
    "Heisman": "Player of the Year", "Maxwell": "Player of the Year", "Camp*": "Most Outstanding Player",
    "Bednarik": "Defensive Player of the Year", "Nagurski": "Defensive Player of the Year",
    "O'Brien": "Nation's best QB", "Walker": "Nation's best RB", "Biletnikoff": "Nation's best WR",
    "Mackey": "Nation's best TE", "Outland": "Nation's best interior lineman", "Rimington": "Nation's best center",
    "Lombardi": "Nation's top lineman or linebacker", "Butkus": "Nation's best LB", "Thorpe": "Nation's best DB",
    "Groza": "Nation's best K", "Guy": "Nation's best P",
}
POSITION_FAMILY = {}
ERRORS = ("#DIV/0!", "#REF!", "#N/A", "#VALUE!", "#NAME?", "#NUM!", "#NULL!")
REPORT = []  # build notes for data/build-report.md


# ---------------------------------------------------------------- helpers
def note(msg):
    REPORT.append(msg)


def is_error(v):
    return isinstance(v, str) and v.strip() in ERRORS


def clean(v):
    """F10: error cells -> None. Strip strings; empty -> None. Integer floats -> int."""
    if v is None or is_error(v):
        return None
    if isinstance(v, str):
        v = re.sub(r"\s+", " ", v).strip()
        return v or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def num(v):
    """F11: numbers stored as text ('3,817') -> number."""
    v = clean(v)
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    s = v.replace(",", "").strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+", s):
        return float(s)
    return None


def ordinal(v):
    """'5th', '9TH' -> 5."""
    v = clean(v)
    if v is None:
        return None
    m = re.match(r"(\d+)", str(v))
    return int(m.group(1)) if m else None


def ordinal_text(n):
    if n is None:
        return None
    suf = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suf}"


def wl(v):
    """'(5-4)' -> (5, 4)."""
    v = clean(v)
    m = re.search(r"(\d+)\s*-\s*(\d+)", str(v)) if v is not None else None
    return (int(m.group(1)), int(m.group(2))) if m else (None, None)


def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def r1(x, d=1):
    return None if x is None else round(x, d)


def div(a, b):
    return None if a is None or not b else a / b


# ---------------------------------------------------------------- config inputs
def load_json(*parts):
    with open(os.path.join(DATA, *parts), encoding="utf-8") as f:
        return json.load(f)


FIXES = {f["id"]: f for f in load_json("fixes.json")["fixes"]}
ALIASES = load_json("aliases.json")
TEAMS_REF = load_json("reference", "teams.json")
SITE = load_json("site-control.json")
for fam, poss in FIXES["F19-position-families"]["set"]["position_family"].items():
    for p in poss:
        POSITION_FAMILY[p] = fam

_alias = {}
for canon, variants in ALIASES["players"].items():
    _alias[canon.lower()] = canon
    for v in variants:
        _alias[re.sub(r"\s+", " ", v).strip().lower()] = canon
_coach_alias = {}
for canon, info in ALIASES["coaches"].items():
    _coach_alias[canon.lower()] = canon
    for v in info["variants"]:
        _coach_alias[v.lower()] = canon
        _coach_alias[re.sub(r"^(coach|hc -)\s+", "", v, flags=re.I).lower()] = canon


def canon(name):
    if name is None:
        return None
    n = re.sub(r"\s+", " ", str(name).replace("*", "")).strip(" -")
    return _alias.get(n.lower(), n)


def canon_coach(name):
    n = re.sub(r"\s+", " ", str(name)).strip()
    return _coach_alias.get(n.lower(), n)


CLASS_RE = re.compile(r"^(FR|SO|JR|SR)\s*(\(\s*(RS|SR)\s*\)?)?$", re.I)


def parse_class(v):
    """F18: 'SR (RS)', 'JR (RS', 'SR ' -> {'year':'SR','rs':True,'label':'SR (RS)'}."""
    v = clean(v)
    if v is None:
        return None
    m = CLASS_RE.match(str(v).strip())
    if not m:
        return None
    year, rs = m.group(1).upper(), bool(m.group(3))
    return {"year": year, "rs": rs, "label": f"{year} (RS)" if rs else year}


LEADER_RE = re.compile(r"^(.*?)(?:\s*-\s*|\s+)((?:FR|SO|JR|SR)\s*(?:\(\s*(?:RS|SR)\s*\)?)?|ROLB)\s*$", re.I)


def split_leader(v):
    """'Connor Bazelak - SR (RS)' -> ('Connor Bazelak', 'SR (RS)'); 'Parker Navarro' -> (name, None)."""
    v = clean(v)
    if v is None:
        return None, None
    m = LEADER_RE.match(v)
    if m:
        return m.group(1).strip(" -"), m.group(2).strip()
    return v.strip(" -"), None


# ---------------------------------------------------------------- workbook access
def newest_source():
    files = sorted(glob.glob(os.path.join(ROOT, "source", "*.xlsx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    if not files:
        sys.exit("No .xlsx in source/")
    return files[-1]


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title

    def v(self, col, row):
        c = CI(col) if isinstance(col, str) else col
        return clean(self.ws.cell(row, c).value)

    def raw(self, col, row):
        c = CI(col) if isinstance(col, str) else col
        return self.ws.cell(row, c).value


# ---------------------------------------------------------------- parsers
def parse_coach_sheets(wb):
    seasons, coaches = [], []
    for prog, name in COACH_SHEETS.items():
        sh = Sheet(wb[name])
        hdr = {}
        for c in range(1, sh.ws.max_column + 1):
            h = sh.v(c, 2)
            if isinstance(h, str):
                hdr.setdefault(h.upper(), c)
        wcol = hdr.get("W") or hdr.get("WIN")
        lcol = hdr.get("L") or hdr.get("LOSS")
        leaders = [("passing", "PASSING LEADER"), ("rushing", "RUSHING LEADER"), ("receiving", "RECEIVING LEADER"),
                   ("tackles", "TACKLE LEADER"), ("sacks", "SACK LEADER"), ("interceptions", "INT LEADER")]
        rank_label = "OFF RNK" if "OFF RNK" in hdr else "PPG RANK"
        for row in range(3, sh.ws.max_row + 1):
            year = num(sh.v("B", row))
            if not isinstance(year, int) or year < 2000:
                continue
            w, l = num(sh.v(wcol, row)), num(sh.v(lcol, row))
            if w is None or l is None:
                continue  # future season, not played yet
            conf_w = conf_l = conf_rank = None
            for c in range(lcol + 1, hdr["PASSING LEADER"]):  # conf record / rank sit between L and the leaders
                val = sh.v(c, row)
                if val is None:
                    continue
                if isinstance(val, str) and val.startswith("("):
                    conf_w, conf_l = wl(val)
                elif ordinal(val) is not None and conf_rank is None:
                    conf_rank = ordinal(val)
            rec = {
                "id": f"{year}-{prog}", "program": prog, "season": year, "coach": canon_coach(f"Coach {name.replace('Coach ', '')}".replace("Coach ", "")),
                "w": w, "l": l, "win_pct": r1(div(w, w + l), 3),
                "conf_w": conf_w, "conf_l": conf_l, "conf_rank": conf_rank, "conf_rank_text": ordinal_text(conf_rank),
                "leaders": {},
                "team": {
                    "pass_yds": num(sh.v(hdr["PASS YDS"], row)), "rush_yds": num(sh.v(hdr["RUSH YDS"], row)),
                    "ppg": num(sh.v(hdr["PPG"], row)),
                    "ppg_rank" if rank_label == "PPG RANK" else "offense_rank": ordinal(sh.v(hdr[rank_label], row)),
                },
            }
            for key, label in leaders:
                c = hdr[label]
                nm, cls = split_leader(sh.v(c, row))
                if nm is None:
                    continue
                entry = {"player": canon(nm), "value": num(sh.v(c + 1, row))}
                if cls:
                    if cls.upper() == "ROLB" or cls.upper().startswith("SR (SR"):
                        entry["class"] = None  # resolved by F18 overrides below
                        entry["class_raw"] = cls
                    else:
                        pc = parse_class(cls)
                        entry["class"] = pc["label"] if pc else None
                rec["leaders"][key] = entry
            seasons.append(rec)
        coaches.append({"program": prog, "role": "HC", "name": canon_coach(name.replace("Coach ", "").strip())})
    # F18 overrides
    for o in FIXES["F18-class-labels"]["set"]["overrides"]:
        for s in seasons:
            if s["program"] == "CLEV" and s["season"] == o["season"]:
                for e in s["leaders"].values():
                    if e["player"] == canon(o["player"]) and e.get("class") is None:
                        e["class"] = o["class"]
                        e.pop("class_raw", None)
    for s in seasons:
        s["coach"] = next(c["name"] for c in coaches if c["program"] == s["program"])
    note(f"Coach sheets: {len(seasons)} program-seasons ({Counter(s['program'] for s in seasons)})")
    return seasons, coaches


def parse_schedule(wb):
    """CLEV regular season, conference title games and bowls. Playoff rows come from the Playoff Tree (F05)."""
    sh = Sheet(wb["Schedule"])
    starts = [c for c in range(1, sh.ws.max_column + 1) if sh.v(c, 3) == "Game"]
    games = []
    for i, st in enumerate(starts):
        season = 2024 + i
        cols = {}
        for c in range(st, st + 8):
            h = sh.v(c, 3)
            if h and h not in cols:
                cols[h] = c
        seq = 0
        for row in range(5, sh.ws.max_row + 1):
            label = sh.v(st, row)
            team, opp = sh.v(cols["Team"], row), sh.v(cols["Opp"], row)
            if label is None or team is None:
                continue
            score_raw, result_raw = sh.v(cols["Final Score"], row), sh.v(cols["Score"], row)
            label_s = str(label).strip()
            if isinstance(label, int) or label_s.isdigit():
                kind, name = "regular", None
            elif label_s.upper().startswith("B12-CH"):
                kind, name = "conference_championship", "Big 12 Championship"
            elif re.match(r"^(RD\s*\d|NC)$", label_s, re.I):
                continue  # F05: playoff games come from the Playoff Tree
            else:
                kind = "bowl"
                name = re.sub(r"\s*(BWL\.?|Bwl)\s*$", " Bowl", label_s, flags=re.I).strip()
                name = " ".join(w if w.isupper() and len(w) <= 3 else w.capitalize() for w in name.split())
            # F02 override
            f02 = FIXES["F02-2034-big12-title-game"]
            if season == f02["match"]["season"] and kind == "conference_championship":
                if team == f02["match"]["seen_value"] or opp == f02["match"]["seen_value"]:
                    s = f02["set"]
                    team, opp, score_raw, result_raw = "CLEV", s["opponent"], f'{s["opp_score"]}-{s["clev_score"]}', s["result"]
                    note(f"F02 applied: {season} Big 12 title game -> CLEV vs {opp} {s['clev_score']}-{s['opp_score']} L")
            # F08 placeholders, F05 non-CLEV rows
            if opp in (None, "N/A") or result_raw in (None, "N/A"):
                continue
            if "CLEV" not in (team, opp):
                continue
            if team != "CLEV":
                team, opp = "CLEV", team
            # F07 result, F06 score
            res = "W" if str(result_raw).upper() in ("W", "CLEV") else "L"
            m = re.match(r"\s*(\d+)\s*-\s*(\d+)\s*$", str(score_raw))
            if not m:
                note(f"Schedule {season} {label_s}: unreadable score {score_raw!r}")
                continue
            a, b = int(m.group(1)), int(m.group(2))
            hi, lo = max(a, b), min(a, b)
            clev, oppsc = (hi, lo) if res == "W" else (lo, hi)
            ha = sh.v(cols["H/A"], row)
            site = "neutral" if kind != "regular" else ("away" if ha == "@" else "home")
            date = None
            if "Date" in cols and sh.v(cols["Date"], row):
                md = str(sh.v(cols["Date"], row))
                mm = re.match(r"(\d+)/(\d+)", md)
                if mm:
                    mo = int(mm.group(1))
                    date = f"{season if mo >= 7 else season + 1}-{mo:02d}-{int(mm.group(2)):02d}"
            seq += 1
            games.append({
                "season": season, "kind": kind, "week": int(label) if kind == "regular" else None, "name": name,
                "opponent": opp, "site": site, "clev_score": clev, "opp_score": oppsc, "result": res, "date": date,
            })
    return games


def parse_playoffs(path):
    seasons = parse_playoff_tree(path)
    round_names = {"first_round": "First Round", "quarterfinal": "Quarterfinal", "semifinal": "Semifinal",
                   "final": "National Championship"}
    rows, champs = [], []
    order = {"first_round": 1, "quarterfinal": 2, "semifinal": 3, "final": 4}
    for s in seasons:
        for g in s["games"]:
            sa, sb = g["score_a"], g["score_b"]
            # F22: strip markers like '73*'
            if isinstance(sa, str):
                sa = num(re.sub(r"[^\d]", "", sa))
            if isinstance(sb, str):
                sb = num(re.sub(r"[^\d]", "", sb))
            winner = None
            if isinstance(sa, int) and isinstance(sb, int) and sa != sb:
                winner = g["team_a"] if sa > sb else g["team_b"]
            rows.append({"season": s["season"], "round": g["round"], "round_name": round_names[g["round"]],
                         "round_order": order[g["round"]], "team_a": g["team_a"], "score_a": sa,
                         "team_b": g["team_b"], "score_b": sb, "winner": winner})
        fin = [r for r in rows if r["season"] == s["season"] and r["round"] == "final"]
        if fin:
            f = fin[0]
            if f["winner"] != s["champion_cell"]:
                note(f"Playoff Tree {s['season']}: final winner {f['winner']} != champion cell {s['champion_cell']}")
            loser = f["team_b"] if f["winner"] == f["team_a"] else f["team_a"]
            champs.append({"season": s["season"], "champion": s["champion_cell"] or f["winner"], "runner_up": loser,
                           "score": f"{max(f['score_a'], f['score_b'])}-{min(f['score_a'], f['score_b'])}",
                           "byes": s.get("byes")})
    return rows, champs


def parse_conference_championships(wb):
    ws = wb["Conference Champions"]
    out = []
    f03 = {s["cell"]: s["conference"] for s in FIXES["F03-2029-conference-labels-swapped"]["set"]}
    tables = sorted(ws.tables.values(), key=lambda t: openpyxl.utils.cell.range_boundaries(t.ref)[1])
    for t in tables:
        rng = ws[t.ref]
        title = clean(rng[0][0].value)
        m = re.match(r"(\d{4})", str(title or ""))
        season = int(m.group(1)) if m else 2024
        for row in rng[1:]:
            conf_raw = clean(row[0].value)
            if row[0].coordinate in f03:
                conf_raw = f03[row[0].coordinate]
            a, b, score, champ = clean(row[1].value), clean(row[3].value), clean(row[4].value), clean(row[5].value)
            if not (a and b and champ):
                continue
            mm = re.match(r"\s*(\d+)\s*-\s*(\d+)", str(score))
            hi, lo = (max(int(mm.group(1)), int(mm.group(2))), min(int(mm.group(1)), int(mm.group(2)))) if mm else (None, None)
            loser = b if champ == a else a
            out.append({"season": season, "conference": CONF_NAMES.get(conf_raw, conf_raw), "champion": champ,
                        "runner_up": loser, "champion_score": hi, "runner_up_score": lo})
    note(f"Conference title games: {len(out)} across {len({o['season'] for o in out})} seasons")
    return out


def parse_standings(wb):
    sh = Sheet(wb["Standings"])
    out, natl = [], {}
    f16 = FIXES["F16-2034-mac-standings-emu"]
    for i in range(0, 40):
        col = 3 + i
        season = num(sh.v(col, 2))
        if not isinstance(season, int):
            continue
        champ, ru = sh.v(col, 3), sh.v(col, 4)
        if champ:
            natl[season] = (champ, ru)
        for k, conf in enumerate(STANDINGS_CONFS):
            top = 6 + 5 * k
            for place in range(4):
                cell = f"{CL(col)}{top + place}"
                team = sh.v(col, top + place)
                if cell == f16["match"]["cell"] and team == f16["match"]["seen_value"]:
                    team = f16["set"]["team"]
                if team:
                    out.append({"season": season, "conference": conf, "place": place + 1, "team": team})
    return out, natl


def parse_ap(wb):
    sh = Sheet(wb["AP Top 25"])
    out = []
    for i in range(0, 40):
        col = 3 + i
        season = num(sh.v(col, 2))
        if not isinstance(season, int):
            continue
        for r in range(3, 28):
            t = sh.v(col, r)
            if t:
                out.append({"season": season, "rank": r - 2, "team": t})
    return out


def parse_awards(wb):
    sh = Sheet(wb["Awards"])
    awards = []
    cols = [(c, sh.v(c, 2)) for c in range(1, sh.ws.max_column + 1) if isinstance(sh.v(c, 2), str)]
    for row in range(3, sh.ws.max_row + 1):
        season = num(sh.v("A", row))
        if not isinstance(season, int):
            continue
        for c, award in cols:
            winner = sh.v(c, row)
            if not winner:
                continue
            key = award.rstrip("*")
            awards.append({"season": season, "award": key, "award_note": AWARDS_INFO.get(award),
                           "player": canon(winner), "school": sh.v(c - 2, row)})
    return awards


def parse_all_americans(wb):
    sh = Sheet(wb["All-Americans"])
    out = []
    unit = None
    team_cols = [c for c in range(1, sh.ws.max_column + 1) if sh.v(c, 3) == "Team"]
    for c in team_cols:
        season = num(sh.v(c, 2))
        if not isinstance(season, int):
            continue
        unit = None
        for row in range(4, 32):
            lab = sh.v("A", row)
            if lab in ("Offense", "Defense", "ST"):
                unit = {"ST": "Special Teams"}.get(lab, lab)
            player, pos, school = sh.v(c + 2, row), sh.v(c + 3, row), sh.v(c, row)
            if player and school:
                out.append({"season": season, "unit": unit, "position": pos, "player": canon(player), "school": school,
                            "team": "First Team"})
    return out


def parse_recruiting(wb):
    sh = Sheet(wb["Recruiting"])
    recruits = []
    starts = [c for c in range(1, sh.ws.max_column + 1) if sh.v(c, 3) == "Nat. Rank"]
    f14_cycle = 2026
    for st in starts:
        cycle = num(sh.v(st, 2))
        if not isinstance(cycle, int):
            continue
        hdr = {}
        for c in range(st, st + 7):
            h = sh.v(c, 3)
            if h and h not in hdr:
                hdr[h] = c
        for row in range(4, 38):
            player = sh.v(hdr["Player"], row)
            if not player or player == "Recruiting" or not re.search(r"[a-z]", str(player)):
                continue
            stars = num(sh.v(hdr["Stars"], row)) if "Stars" in hdr else None
            if cycle == f14_cycle:
                stars = None  # F14
            pos = sh.v(hdr["POS"], row)
            recruits.append({"cycle": cycle, "arrival_season": cycle + 1, "program": "CLEV", "player": canon(player),
                             "natl_rank": num(sh.v(hdr["Nat. Rank"], row)), "position": pos.strip() if pos else None,
                             "stars": stars, "state": sh.v(hdr["State"], row) if "State" in hdr else None,
                             "stage": sh.v(hdr["Stage"], row) if "Stage" in hdr else None})
    # Class summaries (F13): rows 41-43 = CLEV, M-OH, OHIO; 2024 at C-H, 2025 at K-P
    classes = []
    for cycle, c0 in ((2024, "C"), (2025, "K")):
        for i, prog in enumerate(["CLEV", "M-OH", "OHIO"]):
            row = 41 + i
            base = CI(c0)
            counts = [num(sh.v(base + k, row)) for k in range(5)]
            classes.append({"program": prog, "cycle": cycle, "natl_rank": ordinal(sh.v(base + 5, row)),
                            "natl_rank_text": sh.v(base + 5, row),
                            "summary_stars": dict(zip(["5", "4", "3", "2", "1"], counts))})
    by_cycle = defaultdict(list)
    for r in recruits:
        by_cycle[r["cycle"]].append(r)
    for cycle, rs in sorted(by_cycle.items()):
        tracked = [r for r in rs if r["stars"] is not None]
        counts = Counter(int(r["stars"]) for r in tracked)
        row = next((c for c in classes if c["program"] == "CLEV" and c["cycle"] == cycle), None)
        if row is None:
            row = {"program": "CLEV", "cycle": cycle, "natl_rank": None, "natl_rank_text": None}
            classes.append(row)
        row["signees"] = len(rs)
        row["stars"] = {str(k): counts.get(k, 0) for k in (5, 4, 3, 2, 1)} if tracked else None
        row["stars_tracked"] = bool(tracked)
        row["avg_natl_rank"] = r1(sum(r["natl_rank"] for r in rs if r["natl_rank"]) / max(1, len([r for r in rs if r["natl_rank"]])), 0)
    return recruits, classes


def parse_graduates(wb):
    sh = Sheet(wb["Graduates"])
    out = []
    for c in range(1, sh.ws.max_column + 1):
        if sh.v(c, 4) != "Class":
            continue
        title = sh.v(c, 3)
        m = re.match(r"(\d{4})", str(title or ""))
        if not m:
            continue
        season = int(m.group(1))
        for row in range(5, sh.ws.max_row + 1):
            player = sh.v(c + 1, row)
            if not player or not re.search(r"[a-z]", str(player)) or player in ("Graduates", "Player"):
                continue
            drafted_raw = sh.v(c + 4, row)
            rd = sh.v(c + 5, row)
            drafted = {"YES": True, "NO": False}.get(str(drafted_raw).upper()) if drafted_raw else None
            cls = parse_class(sh.v(c, row))
            out.append({"season": season, "program": "CLEV", "player": canon(player), "class": cls["label"] if cls else sh.v(c, row),
                        "position": sh.v(c + 2, row), "overall": num(sh.v(c + 3, row)), "drafted": drafted,
                        "draft_round": ordinal(rd) if drafted else None,
                        "outcome": "cut" if str(rd).upper() == "CUT" else ("drafted" if drafted else ("graduated" if drafted is False else "pending"))})
    return out


def parse_transfers(wb):
    sh = Sheet(wb["Transfers"])
    out = []
    for c in range(1, sh.ws.max_column + 1):
        if sh.v(c, 3) != "Class":
            continue
        cycle = num(sh.v(c, 2))
        if not isinstance(cycle, int):
            continue
        for row in range(4, 12):
            player = sh.v(c + 1, row)
            if not player or not re.search(r"[a-z]", str(player)) or player == "Transfers":
                continue
            raw_cls = sh.v(c, row)
            cls = parse_class(raw_cls)
            if cycle == 2025 and isinstance(raw_cls, int):
                cls = None  # F15: numbers in the Class column, not recorded
            out.append({"cycle": cycle, "arrival_season": cycle + 1, "player": canon(player),
                        "class": cls["label"] if cls else None, "position": sh.v(c + 2, row),
                        "to": sh.v(c + 3, row), "from": sh.v(c + 4, row)})
    return out


DC_LABELS = {"WR", "LT", "LG", "C", "RG", "RT", "TE", "RB", "QB", "FB", "CB 1", "CB 2", "CB 3", "LE", "DT", "RE",
             "FS", "LOLB", "MLB", "ROLB", "SS", "K", "P", "COACHES", "EDGE", "DT 1", "DT 2", "CB"}


def parse_depth_chart(wb, current_season, departing_names):
    sh = Sheet(wb["Depth Chart"])
    years = {c: num(sh.v(c, 2)) for c in range(1, sh.ws.max_column + 1) if isinstance(num(sh.v(c, 2)), int)}
    cols = sorted(years)
    chart, staff, loose_notes, season_notes = [], [], [], []
    for i, st in enumerate(cols):
        season = years[st]
        if season > current_season + 1:
            continue  # F17: blocks after next season are leftovers
        end = cols[i + 1] if i + 1 < len(cols) else st + 12
        rated = season <= 2032
        projected = season == current_season + 1
        for col in range(st - 1, end - 1):
            label, depth = None, 0
            for row in range(3, 29):
                val = sh.v(col, row)
                if val is None:
                    continue
                sval = str(val).strip()
                if sval.upper() in DC_LABELS:
                    label, depth = sval.upper().replace("CB ", "CB"), 0
                    continue
                if label is None or sval.startswith(("Offense", "Bolded", "Italic", "* =")) or sval == "_":
                    continue
                if label == "COACHES":
                    m = re.match(r"^(HC|OC|DC)\s*-\s*(.*)$", sval)
                    if m and m.group(2).strip():
                        nm = m.group(2).strip()
                        stars = len(nm) - len(nm.rstrip("*"))
                        staff.append({"season": season, "role": m.group(1), "name_raw": nm,
                                      "name": canon_coach(nm.rstrip("*").strip()), "footnote_marks": stars})
                    continue
                if len(re.sub(r"\s*-?\s*\d{2}\s*$", "", sval).split()) > 5 or "," in sval:
                    loose_notes.append({"season": season, "text": sval.lstrip("*").strip(),
                                        "marks": len(sval) - len(sval.lstrip("*"))})
                    continue  # a staff note typed into the chart, not a player
                m = re.match(r"^(.*?)\s*-?\s*(\d{2})\s*$", sval)
                name, rating = (m.group(1), int(m.group(2))) if m else (sval, None)
                transfer_mark = "*" in name
                name = canon(name)
                if not re.search(r"[a-z]", name):
                    continue
                depth += 1
                chart.append({"season": season, "program": "CLEV", "slot": label, "depth": depth, "player": name,
                              "rating": rating if rated else None, "transfer_mark": transfer_mark,
                              "projected": projected,
                              "departing": projected and name in departing_names})
        # staff footnotes (rows 29-31)
        notes = []
        for col in range(st - 1, end - 1):
            for row in range(29, 32):
                val = sh.v(col, row)
                if isinstance(val, str) and val.startswith("*") and not val.startswith("* ="):
                    notes.append(val)
        notes = [{"text": n.lstrip("*").strip(), "marks": len(n) - len(n.lstrip("*"))} for n in notes]
        notes += [n for n in loose_notes if n["season"] == season]
        matched = set()
        for s in staff:
            if s["season"] == season:
                s["notes"] = [n["text"] for n in notes if s["footnote_marks"] and n["marks"] == s["footnote_marks"]]
                matched |= {n["text"] for n in notes if s["footnote_marks"] and n["marks"] == s["footnote_marks"]}
        for n in notes:
            if n["text"] not in matched:
                season_notes.append({"season": season, "note": n["text"]})
    return chart, staff, season_notes


def parse_stats(wb, seasons_played):
    rows = []
    kinds = {"COMP": "passing", "CAR": "rushing", "REC": "receiving", "FGM": "kicking"}
    for season in seasons_played:
        name = f"{season} Stats"
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for t in ws.tables.values():
            rng = ws[t.ref]
            hdr = [clean(c.value) for c in rng[0]]
            kind = next((k for key, k in kinds.items() if key in hdr), None)
            if not kind:
                continue
            for r in rng[1:]:
                vals = dict(zip(hdr[1:], [clean(c.value) for c in r[1:]]))
                player = clean(r[0].value)
                if not player or not re.search(r"[a-z]", str(player)):
                    continue
                rec = {"season": season, "program": "CLEV", "category": kind, "player": canon(player),
                       "position": vals.get("POS"), "_cells": {h: f"{c.coordinate}" for h, c in zip(hdr, r)}}
                for h, v in vals.items():
                    if h and h != "POS":
                        rec[h] = num(v)
                rows.append(rec)
    # F04
    f = FIXES["F04-2032-fatafehi-receiving-yards"]
    for r in rows:
        if r["season"] == 2032 and r["category"] == "receiving" and r["player"] == canon(f["match"]["player"]) and r.get("YARDS") == f["match"]["seen_value"]:
            r["YARDS"] = f["set"]["YARDS"]
            note("F04 applied: 2032 Sam Fatafehi receiving yards 2219 -> 219")
    # Normalize field names and recompute averages (F10)
    out = []
    for r in rows:
        k, gp = r["category"], r.get("GP")
        o = {"id": f"{r['season']}-{k}-{slug(r['player'])}", "season": r["season"], "program": "CLEV", "category": k,
             "player": r["player"], "position": r["position"], "gp": gp, "dp": r.get("DP")}
        if k == "passing":
            o.update(yds=r.get("YDS"), comp=r.get("COMP"), att=r.get("ATT"), td=r.get("TD"), int=r.get("INT"),
                     rate=r.get("RATE"), long=r.get("LONG"), sacks=r.get("SACKS"))
            o["comp_pct"] = r1(div(o["comp"], o["att"]) * 100 if div(o["comp"], o["att"]) is not None else None)
            o["ypa"] = r1(div(o["yds"], o["att"]))
            o["ypg"] = r1(div(o["yds"], gp))
        elif k == "rushing":
            o.update(car=r.get("CAR"), yds=r.get("YDS"), td=r.get("TD"), fumb=r.get("FUMB"), bkt=r.get("BKT"),
                     yac=r.get("YAC"), long=r.get("LONG"), plays_20=r.get("20+"))
            o["ypc"] = r1(div(o["yds"], o["car"]))
            o["ypg"] = r1(div(o["yds"], gp))
        elif k == "receiving":
            o.update(rec=r.get("REC"), yds=r.get("YARDS"), td=r.get("TD"), yac=r.get("YAC"), long=r.get("LONG"),
                     drops=r.get("DROPS"))
            o["ypr"] = r1(div(o["yds"], o["rec"]))
            o["ypg"] = r1(div(o["yds"], gp))
        else:
            o.update(fgm=r.get("FGM"), fga=r.get("FGA"), xpm=r.get("XPM"), xpa=r.get("XPA"), long=r.get("LONG"),
                     fg_0_29=[r.get("FGM29"), r.get("FGA29")], fg_30_39=[r.get("FGM39"), r.get("FGA39")],
                     fg_40_49=[r.get("FGM49"), r.get("FGA49")], fg_50_plus=[r.get("FGM50+"), r.get("FGA50+")])
            o["fg_pct"] = r1(div(o["fgm"], o["fga"]) * 100 if div(o["fgm"], o["fga"]) is not None else None)
            o["xp_pct"] = r1(div(o["xpm"], o["xpa"]) * 100 if div(o["xpm"], o["xpa"]) is not None else None)
        out.append(o)
    return out


def parse_records(wb):
    sh = Sheet(wb["RecordsHall of Fame"])
    out = []
    for row in range(2, sh.ws.max_row + 1):
        player = sh.v("A", row)
        if not player:
            continue
        mag = str(sh.v("C", row) or "").upper()
        mag = {"CLEVELAND STATE": "Cleveland State", "BIG 12": "Big 12", "NCAA": "NCAA"}.get(mag, mag.title())
        out.append({"player": canon(player), "position": sh.v("B", row), "magnitude": mag,
                    "duration": sh.v("D", row), "type": sh.v("E", row), "value": num(sh.v("F", row)),
                    "year_set": str(sh.v("G", row) or "")})
    return out


def computed_bests(player_seasons, top=10):
    """F12: CLEV season and career bests from the stats tabs."""
    stats = {
        "passing": [("yds", "Passing Yards"), ("td", "Passing TD"), ("comp", "Completions")],
        "rushing": [("yds", "Rushing Yards"), ("td", "Rushing TD")],
        "receiving": [("rec", "Receptions"), ("yds", "Receiving Yards"), ("td", "Receiving TD")],
        "kicking": [("fgm", "Field Goals Made")],
    }
    out = []
    for cat, fields in stats.items():
        rows = [r for r in player_seasons if r["category"] == cat]
        for f, label in fields:
            season = sorted([r for r in rows if r.get(f) is not None], key=lambda r: -r[f])[:top]
            out.append({"scope": "season", "category": cat, "stat": f, "label": label,
                        "leaders": [{"player": r["player"], "season": r["season"], "value": r[f]} for r in season]})
            car = defaultdict(lambda: {"value": 0, "seasons": []})
            for r in rows:
                if r.get(f) is not None:
                    car[r["player"]]["value"] += r[f]
                    car[r["player"]]["seasons"].append(r["season"])
            career = sorted(car.items(), key=lambda kv: -kv[1]["value"])[:top]
            out.append({"scope": "career", "category": cat, "stat": f, "label": label,
                        "leaders": [{"player": p, "value": d["value"], "first": min(d["seasons"]), "last": max(d["seasons"])}
                                    for p, d in career]})
    return out


def team_colors(wb):
    ws = wb["Instructions"]
    out = {}
    for row in ws.iter_rows(min_row=24):
        for c in row:
            if isinstance(c.value, str) and 1 < len(c.value.strip()) <= 5:
                fill = c.fill.fgColor if c.fill and c.fill.fill_type else None
                font = c.font.color if c.font else None
                hexify = lambda col: ("#" + col.rgb[-6:]) if col is not None and col.type == "rgb" and isinstance(col.rgb, str) else None
                out[c.value.strip()] = {"primary": hexify(fill), "text": hexify(font)}
    return out


def snapshot(wb):
    os.makedirs(SNAP, exist_ok=True)
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    v = c.value
                    cells[c.coordinate] = v if isinstance(v, (int, float, str)) else str(v)
        fn = slug(ws.title) + ".json"
        with open(os.path.join(SNAP, fn), "w", encoding="utf-8") as f:
            json.dump(cells, f, ensure_ascii=False, separators=(",", ":"))


# ---------------------------------------------------------------- assemble
def build(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    wb_styles = openpyxl.load_workbook(path)  # colors live in styles
    sha = hashlib.sha256(open(path, "rb").read()).hexdigest()

    seasons, hcs = parse_coach_sheets(wb)
    current = max(s["season"] for s in seasons)
    if current != SITE["current_season"]:
        note(f"WARNING: tracker's latest season is {current} but site-control.json says {SITE['current_season']}")
    seasons_played = sorted({s["season"] for s in seasons})

    sched = parse_schedule(wb)
    playoffs, natl_from_tree = parse_playoffs(path)
    ccg = parse_conference_championships(wb)
    standings, natl_from_standings = parse_standings(wb)
    ap = parse_ap(wb)
    awards = parse_awards(wb)
    aas = parse_all_americans(wb)
    recruits, rclasses = parse_recruiting(wb)
    departures = parse_graduates(wb)
    transfers = parse_transfers(wb)
    latest_departing = {d["player"] for d in departures if d["season"] == current}
    depth, staff, staff_notes = parse_depth_chart(wb, current, latest_departing)
    pstats = parse_stats(wb, seasons_played)
    records = parse_records(wb)
    bests = computed_bests(pstats)
    colors = team_colors(wb_styles)

    # National championships: tree is the source (F05); standings cross-checked
    natl = []
    for n in natl_from_tree:
        st = natl_from_standings.get(n["season"])
        if st and (st[0] != n["champion"] or st[1] != n["runner_up"]):
            note(f"National title {n['season']}: tree {n['champion']}/{n['runner_up']} vs standings {st}")
        natl.append(n)

    # CLEV playoff games from the tree (F05)
    for p in playoffs:
        if "CLEV" in (p["team_a"], p["team_b"]) and p["season"] >= 2024:
            me_a = p["team_a"] == "CLEV"
            clev, opp = (p["score_a"], p["score_b"]) if me_a else (p["score_b"], p["score_a"])
            sched.append({"season": p["season"], "kind": "playoff", "week": None,
                          "name": "CFP " + p["round_name"] if p["round"] != "final" else "National Championship",
                          "round": p["round"], "opponent": p["team_b"] if me_a else p["team_a"], "site": "neutral",
                          "clev_score": clev, "opp_score": opp, "result": "W" if clev > opp else "L", "date": None})
    kind_order = {"regular": 0, "conference_championship": 1, "bowl": 2, "playoff": 3}
    round_order = {"first_round": 1, "quarterfinal": 2, "semifinal": 3, "final": 4}
    sched.sort(key=lambda g: (g["season"], kind_order[g["kind"]], g["week"] or 0, round_order.get(g.get("round"), 0)))
    games = []
    seq = Counter()
    for g in sched:
        seq[g["season"]] += 1
        g = {"id": f"{g['season']}-CLEV-{seq[g['season']]:02d}", "program": "CLEV", "game_no": seq[g["season"]], **g}
        games.append(g)

    # Season postseason summaries + AP rank + titles
    ap_rank = {(a["season"], a["team"]): a["rank"] for a in ap}
    ccg_by = defaultdict(list)
    for c in ccg:
        ccg_by[c["season"]].append(c)
    for s in seasons:
        p, y = s["program"], s["season"]
        s["ap_final_rank"] = ap_rank.get((y, p))
        nat = next((n for n in natl if n["season"] == y), None)
        s["national_champion"] = bool(nat and nat["champion"] == p)
        s["national_runner_up"] = bool(nat and nat["runner_up"] == p)
        b12 = next((c for c in ccg_by[y] if c["conference"] == "Big 12"), None)
        s["conference_champion"] = bool(b12 and b12["champion"] == p)
        s["conference_title_game"] = bool(b12 and p in (b12["champion"], b12["runner_up"]))
        pg = [x for x in playoffs if x["season"] == y and p in (x["team_a"], x["team_b"])]
        if pg:
            last = max(pg, key=lambda x: x["round_order"])
            s["playoff"] = {"games": len(pg), "wins": sum(1 for x in pg if x["winner"] == p),
                            "exit_round": last["round_name"], "won_title": last["round"] == "final" and last["winner"] == p}
        else:
            s["playoff"] = None
        if p == "CLEV":
            bowl = next((g for g in games if g["season"] == y and g["kind"] == "bowl"), None)
            s["bowl"] = {"name": bowl["name"], "opponent": bowl["opponent"], "result": bowl["result"],
                         "score": f"{bowl['clev_score']}-{bowl['opp_score']}"} if bowl else None
            gw = sum(1 for g in games if g["season"] == y and g["result"] == "W")
            gl = sum(1 for g in games if g["season"] == y and g["result"] == "L")
            s["games_record_check"] = f"{gw}-{gl}"
            if (gw, gl) != (s["w"], s["l"]):
                note(f"CHECK: CLEV {y} games add up to {gw}-{gl}, coach sheet says {s['w']}-{s['l']}")
        # attach leader classes for CLEV leaders from stats-free sources stay as entered

    # Programs (F01 career totals)
    programs = []
    for p in PROGRAMS:
        ss = [s for s in seasons if s["program"] == p]
        w, l = sum(s["w"] for s in ss), sum(s["l"] for s in ss)
        programs.append({
            "id": p, "school": SCHOOLS[p], "nickname": NICKNAMES[p], "coach": next(h["name"] for h in hcs if h["program"] == p),
            "conference": "Big 12", "logo": f"assets/logos/{p}.png", "colors": colors.get(p),
            "first_season": min(s["season"] for s in ss), "last_season": max(s["season"] for s in ss),
            "career": {"w": w, "l": l, "win_pct": r1(div(w, w + l), 3)},
            "national_titles": [n["season"] for n in natl if n["champion"] == p and n["season"] >= 2024],
            "national_runner_up": [n["season"] for n in natl if n["runner_up"] == p and n["season"] >= 2024],
            "big12_titles": [c["season"] for c in ccg if c["conference"] == "Big 12" and c["champion"] == p],
            "playoff_appearances": sorted({x["season"] for x in playoffs if p in (x["team_a"], x["team_b"])}),
            "ap_top25_finishes": sum(1 for s in ss if s["ap_final_rank"]),
        })

    # Coaches table (HC per program-season + CLEV staff)
    coaches = []
    hc_notes = {st["season"]: st.get("notes", []) for st in staff if st["role"] == "HC"}
    for s in seasons:
        row = {"program": s["program"], "season": s["season"], "role": "HC", "name": s["coach"]}
        if s["program"] == "CLEV" and hc_notes.get(s["season"]):
            row["notes"] = hc_notes[s["season"]]
        coaches.append(row)
    for st in staff:
        if st["role"] != "HC" and st["season"] <= current:
            coaches.append({"program": "CLEV", "season": st["season"], "role": st["role"], "name": st["name"],
                            "notes": st.get("notes", [])})
    for n in staff_notes:
        if n["season"] <= current:
            coaches.append({"program": "CLEV", "season": n["season"], "role": "note", "name": None, "notes": [n["note"]]})

    # Players registry
    players = {}

    def P(name, team=None):
        pid = slug(name)
        p = players.get(pid)
        if p is None:
            p = players[pid] = {"id": pid, "name": name, "teams": [], "positions": [], "seasons": [],
                                "ratings": {}, "recruit": None, "transfer": None, "departure": None,
                                "awards": [], "all_american": [], "records": [], "leader_seasons": []}
        if team and team not in p["teams"]:
            p["teams"].append(team)
        return p

    def add_pos(p, pos):
        if pos and pos not in p["positions"]:
            p["positions"].append(pos)

    def add_season(p, y):
        if y not in p["seasons"]:
            p["seasons"].append(y)

    for r in pstats:
        p = P(r["player"], "CLEV"); add_pos(p, r["position"]); add_season(p, r["season"])
    for d in depth:
        if d["projected"]:
            continue
        p = P(d["player"], "CLEV"); add_season(p, d["season"])
        if d["rating"]:
            p["ratings"][str(d["season"])] = max(d["rating"], p["ratings"].get(str(d["season"]), 0))
        slot = re.sub(r"\d", "", d["slot"]).strip()
        if slot not in ("COACHES",):
            add_pos(p, slot)
    for r in recruits:
        p = P(r["player"], "CLEV"); add_pos(p, r["position"])
        p["recruit"] = {k: r[k] for k in ("cycle", "arrival_season", "natl_rank", "stars", "state", "position")}
        p["home_state"] = r["state"]
    for t in transfers:
        p = P(t["player"], t["to"]); add_pos(p, t["position"])
        p["transfer"] = {"cycle": t["cycle"], "arrival_season": t["arrival_season"], "from": t["from"], "to": t["to"], "class": t["class"]}
        if t["from"] and t["from"] not in p["teams"]:
            p["teams"].insert(0, t["from"])
    for d in departures:
        p = P(d["player"], "CLEV"); add_pos(p, d["position"]); add_season(p, d["season"])
        p["departure"] = {k: d[k] for k in ("season", "class", "overall", "drafted", "draft_round", "outcome")}
    for s in seasons:
        for cat, e in s["leaders"].items():
            p = P(e["player"], s["program"]); add_season(p, s["season"])
            p["leader_seasons"].append({"season": s["season"], "program": s["program"], "category": cat, "value": e["value"], "class": e.get("class")})
    for a in awards:
        p = P(a["player"], a["school"])
        p["awards"].append({"season": a["season"], "award": a["award"], "school": a["school"]})
    for a in aas:
        p = P(a["player"], a["school"]); add_pos(p, a["position"])
        p["all_american"].append({"season": a["season"], "position": a["position"], "school": a["school"]})
    for rec in records:
        p = P(rec["player"], "CLEV")
        p["records"].append({k: rec[k] for k in ("magnitude", "duration", "type", "value", "year_set")})
    for p in players.values():
        p["seasons"].sort()
        p["position_family"] = Counter(POSITION_FAMILY.get(x, "Other") for x in p["positions"]).most_common(1)[0][0] if p["positions"] else None
        peaks = [v for v in p["ratings"].values()] + ([p["departure"]["overall"]] if p["departure"] and p["departure"]["overall"] else [])
        p["peak_rating"] = max(peaks) if peaks else None
        p["program"] = next((t for t in p["teams"] if t in PROGRAMS), p["teams"][0] if p["teams"] else None)
        p.setdefault("home_state", None)
    confirmed = ALIASES.get("confirmed_same_person", {})
    for name, info in confirmed.items():
        if slug(name) in players:
            players[slug(name)]["career_note"] = info["note"]
    multi = [p for p in players.values() if len([t for t in p["teams"] if t not in (None,)]) > 1 and not p["transfer"] and p["name"] not in confirmed]
    for p in multi:
        note(f"Player listed with more than one school (check it's one person): {p['name']} {p['teams']}")

    # Team table
    teams = []
    names = TEAMS_REF["teams"]
    seen_abbr = set(names) | set(colors)
    for abbr in sorted(seen_abbr):
        if abbr in ("W", "L"):
            continue
        logo = os.path.join(ROOT, "site", "assets", "logos", f"{abbr}.png")
        teams.append({"id": abbr, "school": names.get(abbr), "fcs": abbr in TEAMS_REF["fcs"],
                      "colors": colors.get(abbr), "logo": f"assets/logos/{abbr}.png" if os.path.exists(logo) else None})
    used = {g["opponent"] for g in games} | {x for p in playoffs for x in (p["team_a"], p["team_b"])} | {s["team"] for s in standings} | {a["team"] for a in ap} | {a["school"] for a in awards} | {a["school"] for a in aas}
    missing = sorted(t for t in used if t and t not in names)
    if missing:
        note(f"Team abbreviations with no school name in reference/teams.json: {missing}")

    # Strip helper fields
    for r in pstats:
        r.pop("_cells", None)

    tables = {
        "site_control": {**{k: SITE[k] for k in ("current_season", "current_phase", "phase_display", "phases")},
                         "latest_season_in_tracker": current,
                         "last_updated": datetime.date.today().isoformat(),
                         "source_file": os.path.basename(path), "source_sha256": sha},
        "programs": programs, "coaches": coaches, "seasons": seasons, "games": games, "playoffs": playoffs,
        "conference_championships": ccg, "national_championships": natl, "standings": standings, "ap_top25": ap,
        "player_seasons": pstats, "players": sorted(players.values(), key=lambda p: p["id"]), "recruits": recruits,
        "recruiting_classes": sorted(rclasses, key=lambda c: (c["program"], c["cycle"])), "transfers": transfers,
        "departures": departures, "depth_chart": depth, "awards": awards, "all_americans": aas,
        "records_official": records, "records_computed": bests, "teams": teams, "clips": [], "columns": [],
    }
    os.makedirs(OUT, exist_ok=True)
    manifest = {}
    for name, data in tables.items():
        with open(os.path.join(OUT, f"{name}.json"), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        manifest[name] = len(data) if isinstance(data, list) else 1
    snapshot(wb)
    with open(os.path.join(OUT, "_manifest.json"), "w", encoding="utf-8") as f:
        json.dump({"built": datetime.datetime.now().isoformat(timespec="seconds"), "source_file": os.path.basename(path),
                   "source_sha256": sha, "fixes_applied": sorted(FIXES), "row_counts": manifest}, f, indent=1)
    with open(os.path.join(DATA, "build-report.md"), "w", encoding="utf-8") as f:
        f.write(f"# Build report\n\nSource: `{os.path.basename(path)}`\n\n## Row counts\n\n")
        for k, v in manifest.items():
            f.write(f"- {k}: {v}\n")
        f.write("\n## Notes\n\n" + "\n".join(f"- {n}" for n in REPORT) + "\n")
    print("\n".join(f"{k:26} {v}" for k, v in manifest.items()))
    print("\nNOTES:\n" + "\n".join(REPORT))


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else newest_source())
