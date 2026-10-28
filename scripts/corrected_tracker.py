"""Make a corrected COPY of the tracker: every approved fix (data/fixes.json) and name spelling
(data/aliases.json) applied to the cells, plus a "Valhalla Changes" sheet listing each edit
(sheet, cell, old value, new value, fix id, why) so the owner can copy them into the live sheet.

The original tracker and the dated copies in source/ are never modified. The copy is edited by
Excel itself (through PowerShell), because saving with openpyxl would delete the tracker's
186 logo images. Formulas that depend on a corrected cell are recalculated by Excel.

  py -P scripts/corrected_tracker.py [path/to/tracker.xlsx]       (default: newest source/*.xlsx)

Output: source/corrected/<date> CSU Vikings Tracker (corrected).xlsx (git-ignored, stays private)
        data/corrected-tracker-changes.md (the same change list, for the record)
Fixes that are reading rules rather than cell edits (F05-F15, F17, F19) are listed at the
bottom of the Changes sheet with "No cell change"; F05's implied Schedule edits ARE applied.
"""
import datetime as dt
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TRACKER_NAME = "CSU Vikings Tracker"


def load_fixes():
    d = json.load(open(os.path.join(ROOT, "data", "fixes.json"), encoding="utf-8"))
    return {f["id"].split("-")[0]: f for f in d["fixes"]}


class Changes:
    def __init__(self, wb):
        self.wb = wb
        self.items = {}   # (sheet, cell) -> dict
        self.problems = []

    def set(self, sheet, cell, new, fix, why, seen=None, formula=False):
        old = self.wb[sheet][cell].value
        key = (sheet, cell)
        orig = self.items[key]["old"] if key in self.items else old
        if seen is not None and key not in self.items and not same(old, seen):
            self.problems.append(f"{fix} {sheet}!{cell}: expected {seen!r}, found {old!r}; skipped")
            return
        if key in self.items:   # a later fix overrides an earlier one on the same cell
            it = self.items[key]
            it.update(new=new, formula=formula)
            if fix not in it["fix"]:
                it["fix"] += f", {fix}"
                it["why"] += f" Then {why[0].lower() + why[1:]}"
        else:
            self.items[key] = {"sheet": sheet, "cell": cell, "old": orig, "new": new, "fix": fix, "why": why, "formula": formula}


def same(a, b):
    def norm(v):
        if isinstance(v, (int, float)):
            return float(v)
        s = str(v).strip().replace(",", "")
        try:
            return float(s)
        except ValueError:
            return s
    return norm(a) == norm(b)


def build_changes(wb):
    F = load_fixes()
    ch = Changes(wb)

    # F01: career record cells were stale; set the totals computed from the season rows and a live win %.
    careers = {"Coach Josh Dolan": (("F35", 180), ("G35", 29), ("E35", "=F35/(F35+G35)")),
               "Coach Tim Slocombe ": (("G35", 175), ("H35", 34), ("F35", "=G35/(G35+H35)")),
               "Coach Joe Snyderwine": (("G34", 174), ("H34", 32), ("E34", "=G34/(G34+H34)"))}
    for sheet, cells in careers.items():
        for cell, val in cells:
            isf = isinstance(val, str) and val.startswith("=")
            ch.set(sheet, cell, val, "F01", ("Career win % now calculated from the corrected career wins and losses."
                                              if isf else "Career wins/losses were out of date; this is the sum of every season row (2024-2037)."),
                   formula=isf)

    # F02: 2034 Big 12 title game was entered under OHIO
    f = F["F02"]["match"]
    ch.set("Schedule", f["cell_team"], "CLEV", "F02", "2034 Big 12 title game was Cleveland State vs Miami (OH) (lost 42-57), entered under OHIO.", seen=f["seen_value"])

    # F03: 2029 conference labels swapped
    for s in F["F03"]["set"]:
        label = {"MAC": "MAC", "B12": "B12"}[s["conference"]]
        ch.set("Conference Champions", s["cell"], label, "F03", f"2029 labels were swapped: {s['game']} is the {s['conference']} title game.", seen=s["seen_value"])

    # F04
    m = F["F04"]["match"]
    ch.set(m["sheet"], m["cell"], 219, "F04", "Sam Fatafehi 2032 receiving yards: 2219 was a typo for 219 (his yards per catch and per game recalculate).", seen=m["seen_value"])

    # F05: Playoff Tree decides playoff scores and champions -> Schedule cells that disagree
    f05 = [("DC19", "37-27", "'37-3", "2036 first round vs LSU: the Playoff Tree has 37-3."),
           ("DK22", "56-35", "'56-39", "2037 national title vs Ohio: the Playoff Tree has 56-39."),
           ("AJ20", "0-0", "'31-17", "2027 national title (Miami (OH) vs Alabama): the Playoff Tree has Alabama winning 31-17."),
           ("AK20", "M-OH", "BAMA", "2027 national title winner was Alabama (Playoff Tree, Standings and AP poll agree).")]
    for cell, seen, new, why in f05:
        ch.set("Schedule", cell, new, "F05", why, seen=seen)

    # F16, F23
    m = F["F16"]["match"]
    ch.set(m["sheet"], m["cell"], F["F16"]["set"]["team"], "F16", "2034 MAC standings: ECU isn't in the MAC; 2nd place was Eastern Michigan (EMU).", seen=m["seen_value"])
    m = F["F23"]["match"]
    ch.set(m["sheet"], m["cell"], F["F23"]["set"]["school"], "F23", "Ollie Gordon III played for Oklahoma State (OKST), not Oregon State.", seen=m["seen_value"])

    # F18: class labels in the coach sheet
    ch.set("Coach Josh Dolan", "U7", "Dorian Pringle - SR (RS)", "F18", "2028 tackle leader's class was entered as his position (ROLB); 2028 Graduates list has him SR (RS).", seen="Dorian Pringle - ROLB")
    ch.set("Coach Josh Dolan", "X15", "Vonn Ladd - SR (RS)", "F18", "2036 class typo 'SR (SR)'; 2036 Graduates list has him SR (RS).", seen="Vonn Ladd -  SR (SR)")

    # F22: asterisk on the 2035 title game score
    ch.set("CFB Playoff Tree", "H198", 73, "F22", "2035 title game score '73*': the asterisk has no meaning (owner); Miami (OH) 73, Cleveland State 70.", seen="73*")

    # F24-F26
    for s in F["F24"]["set"]:
        ch.set("RecordsHall of Fame", s["cell"], s["value"], "F24", f"{s['player']} career {s['type'].lower()} was entered mid-career; full total {s['why']}.", seen=s["seen_value"])
    for s in F["F25"]["set"]:
        ch.set("2029 Stats", s["cell"], s["value"], "F25", f"{s['player']} 2029 receptions: yards had been pasted into the catches column; real count from the game.", seen=s["seen_value"])
    for s in F["F26"]["set"]:
        ch.set(s["sheet"], s["cell"], "'" + s["value"], "F26", f"2029 Big 12 order (owner decision): Ohio 2nd, Cleveland State 3rd, matching the Standings sheet.", seen=s["seen_value"])

    # F27, F28: records the stats tabs (or the Vikings) beat
    for fid in ("F27", "F28"):
        for s in F[fid]["set"]:
            label = f"{s.get('magnitude', 'Cleveland State')} {s['duration'].lower()} {s['type'].lower()}"
            why = (f"{label} record: {s['player']} {s['value']:,} ({s['year_set']}) beats the entered {s['seen_player']} {s['seen_value']:,}"
                   + (" (from the stats tabs)." if fid == "F27" else " (a Vikings stat line beats it)."))
            c = s["cells"]
            ch.set("RecordsHall of Fame", c["player"], s["player"], fid, why, seen=None)
            ch.set("RecordsHall of Fame", c["position"], s["position"], fid, why, seen=None)
            ch.set("RecordsHall of Fame", c["value"], s["value"], fid, why, seen=s["seen_value"] if (("RecordsHall of Fame", c["value"]) not in ch.items) else None)
            ch.set("RecordsHall of Fame", c["year_set"], "'" + str(s["year_set"]), fid, why, seen=None)

    # Names (aliases.json): only where the text really differs from the approved spelling
    aliases = json.load(open(os.path.join(ROOT, "data", "aliases.json"), encoding="utf-8"))
    var = {v: canon for canon, vs in aliases["players"].items() for v in vs}
    var.update({"Joshua Dolan": "Josh Dolan", "Timmy Slocombe": "Tim Slocombe", "Joseph Snyderwine": "Joe Snyderwine"})
    pats = sorted(var, key=len, reverse=True)
    rx = re.compile(r"(?<![A-Za-z])(" + "|".join(re.escape(p) for p in pats) + r")(?![A-Za-z])", re.I)
    lookup = {p.lower(): var[p] for p in pats}
    skipped_formulas = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if v.startswith("="):
                    if rx.search(v):
                        skipped_formulas.append(f"{ws.title}!{c.coordinate}")
                    continue
                out, pos, changed, fixed = [], 0, False, []
                for m in rx.finditer(v):
                    canon = lookup[m.group(1).lower()]
                    if v[m.start():].startswith(canon):      # already correct (e.g. "RJ Garcia II", "Marcus Sanders Jr.")
                        continue
                    out.append(v[pos:m.start()] + canon)
                    pos = m.end()
                    changed = True
                    fixed.append(f"{m.group(1)} -> {canon}")
                if changed:
                    new = "".join(out) + v[pos:]
                    key = (ws.title, c.coordinate)
                    base = ch.items[key]["new"] if key in ch.items else None
                    if base is not None:
                        continue
                    ch.set(ws.title, c.coordinate, "'" + new if re.fullmatch(r"[\d\s\-/.,:]+", new) else new, "Names",
                           "Approved spelling: " + "; ".join(fixed) + ".")
    return ch, skipped_formulas


RULES = [
    ("F05", "Playoff Tree decides playoff scores and champions; Conference Champions decides other teams' title games; Schedule is used for Vikings games. (The Schedule cells that disagreed are corrected above.)"),
    ("F06", "Scores are read from the Vikings' side (the bigger number goes to the winner); stray spaces ignored."),
    ("F07", "Results read as W/L when a team abbreviation is entered instead."),
    ("F08", "Placeholder rows (N/A, 0-0) are skipped."),
    ("F09", "Each stats row's season comes from the tab name."),
    ("F10", "#DIV/0!, #REF! and #N/A cells are read as blank (they come from formulas on empty rows)."),
    ("F11", "Numbers typed as text ('3,817', '9TH', '(5-4)') are read as numbers."),
    ("F12", "The Records Book shows the Records sheet plus a separate list computed from the stats tabs."),
    ("F13", "Recruiting class summaries: star counts computed from the recruit lists; the 2029-31 summary blocks (copies of 2025) are ignored."),
    ("F14", "2026 recruit stars weren't tracked; shown as 'Not tracked', never estimated."),
    ("F15", "The six 2025 transfers have no class recorded; shown as 'Not recorded'."),
    ("F17", "Depth chart ratings read for 2024-2032; 2033-2037 names only; the 2038 block is the projected chart."),
    ("F19", "Positions kept as entered; the site adds position groups (QB, RB, WR, ...)."),
]


def ps_quote(s):
    return "'" + str(s).replace("'", "''") + "'"


def apply_with_excel(src, dest, items, summary_rows):
    """Copy the tracker and let Excel make the edits and add the Changes sheet."""
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    shutil.copy2(src, dest)
    work = os.path.join(os.path.dirname(dest), "_changes.json")
    json.dump({"items": items, "summary": summary_rows}, open(work, "w", encoding="utf-8"), ensure_ascii=False)
    ps = os.path.join(os.path.dirname(dest), "_apply.ps1")
    script = r"""
$ErrorActionPreference = 'Stop'
$data = Get-Content -Raw -Encoding UTF8 -Path %WORK% | ConvertFrom-Json
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false; $xl.AskToUpdateLinks = $false; $xl.ScreenUpdating = $false
try {
  $wb = $xl.Workbooks.Open(%DEST%, 0, $false)
  foreach ($it in $data.items) {
    $r = $wb.Worksheets.Item($it.sheet).Range($it.cell)
    if ($it.formula) { $r.Formula = [string]$it.new } elseif ($it.new -is [string]) { $r.Value2 = [string]$it.new } else { $r.Value2 = [double]$it.new }
  }
  $last = $wb.Worksheets.Item($wb.Worksheets.Count)
  $ws = $wb.Worksheets.Add([Type]::Missing, $last)
  $ws.Name = 'Valhalla Changes'
  $rows = $data.summary
  $n = $rows.Count; $m = 6
  $arr = New-Object 'object[,]' $n, $m
  for ($i = 0; $i -lt $n; $i++) { for ($j = 0; $j -lt $m; $j++) {
    $v = $rows[$i][$j]
    if ($null -eq $v) { $arr[$i, $j] = '' } elseif ($v -is [string]) { $arr[$i, $j] = [string]$v } else { $arr[$i, $j] = [double]$v }
  } }
  $ws.Range($ws.Cells.Item(1, 1), $ws.Cells.Item($n, $m)).Value2 = $arr
  $ws.Rows.Item(1).Font.Bold = $true; $ws.Rows.Item(3).Font.Bold = $true
  $ws.Range('A1').Font.Size = 14
  $ws.Columns.Item(1).ColumnWidth = 22; $ws.Columns.Item(2).ColumnWidth = 9
  $ws.Columns.Item(3).ColumnWidth = 28; $ws.Columns.Item(4).ColumnWidth = 28
  $ws.Columns.Item(5).ColumnWidth = 12; $ws.Columns.Item(6).ColumnWidth = 90
  $ws.Range('A3:F3').Interior.Color = 3229442  # Valhalla forest (BGR)
  $ws.Range('A3:F3').Font.Color = 16777215
  $ws.Activate(); $xl.ActiveWindow.SplitRow = 3; $xl.ActiveWindow.FreezePanes = $true
  $wb.Worksheets.Item(1).Activate()
  $xl.CalculateFull()
  $wb.Save()
  $wb.Close($false)
} finally {
  $xl.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) | Out-Null
}
'done'
""".replace("%WORK%", ps_quote(work)).replace("%DEST%", ps_quote(dest))
    open(ps, "w", encoding="utf-8-sig").write(script)
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps], capture_output=True, text=True, timeout=600)
    os.remove(ps)
    os.remove(work)
    if r.returncode or "done" not in r.stdout:
        sys.exit("Excel step failed:\n" + r.stdout[-2000:] + r.stderr[-2000:])


def verify(src, dest, items):
    """Every original cell is unchanged except the listed ones, which hold the new values; pictures kept."""
    a = openpyxl.load_workbook(src, data_only=False)
    b = openpyxl.load_workbook(dest, data_only=False)
    want = {(i["sheet"], i["cell"]): i for i in items}
    unexpected, wrong = [], []
    for ws in a.worksheets:
        wb2 = b[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v1, v2 = c.value, wb2[c.coordinate].value
                key = (ws.title, c.coordinate)
                if key in want:
                    exp = want[key]["new"]
                    exp = exp[1:] if isinstance(exp, str) and exp.startswith("'") else exp
                    if want[key]["formula"]:
                        ok = isinstance(v2, str) and v2.replace(" ", "").lower() == str(exp).replace(" ", "").lower()
                    else:
                        ok = same(v2, exp) if not isinstance(exp, str) else v2 == exp
                    if not ok:
                        wrong.append(f"{ws.title}!{c.coordinate}: wanted {exp!r}, got {v2!r}")
                    continue
                if isinstance(v1, str) and v1.startswith("="):
                    if not (isinstance(v2, str) and v2.replace(" ", "").lower() == v1.replace(" ", "").lower()):
                        unexpected.append(f"{ws.title}!{c.coordinate}: formula {v1!r} -> {v2!r}")
                elif not same(v1, v2) if (v1 is not None and v2 is not None) else (v1 not in (None, "") or v2 not in (None, "")):
                    unexpected.append(f"{ws.title}!{c.coordinate}: {v1!r} -> {v2!r}")
    media = lambda p: sum(1 for n in zipfile.ZipFile(p).namelist() if n.startswith("xl/media/"))
    return unexpected, wrong, media(src), media(dest), b.sheetnames


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(ROOT, "source", "*.xlsx")))[-1]
    wb = openpyxl.load_workbook(src, data_only=False)
    ch, skipped = build_changes(wb)
    if ch.problems:
        print("Problems (nothing written):"); [print("  " + p) for p in ch.problems]; sys.exit(1)
    order = {s: i for i, s in enumerate(wb.sheetnames)}
    def cellkey(c):
        m = re.match(r"([A-Z]+)(\d+)", c)
        col = 0
        for ch_ in m.group(1):
            col = col * 26 + ord(ch_) - 64
        return (int(m.group(2)), col)
    items = sorted(ch.items.values(), key=lambda i: (order[i["sheet"]], cellkey(i["cell"])))
    shown = lambda v: "" if v is None else (v[1:] if isinstance(v, str) and v.startswith("'") else v)
    today = dt.date.today().isoformat()
    rows = [[f"Valhalla Changes: {len(items)} cell edits from the approved fixes and name spellings ({today})", "", "", "", "", ""],
            ["Copy each New value into the same sheet and cell of your live tracker. Fix ids match data/fixes.json in the Archive.", "", "", "", "", ""],
            ["Sheet", "Cell", "Old value", "New value", "Fix", "Why"]]
    for i in items:
        old, new = shown(i["old"]), shown(i["new"])
        rows.append([i["sheet"], i["cell"], "'" + str(old) if isinstance(old, str) and not str(old).startswith("=") else ("'" + old if isinstance(old, str) else old),
                     "'" + str(new) if isinstance(new, str) else new, i["fix"], i["why"]])
    rows.append(["", "", "", "", "", ""])
    rows.append(["Approved reading rules (no cell change needed)", "", "", "", "", ""])
    for fid, text in RULES:
        rows.append(["", "", "", "", fid, text])
    rows.append(["", "", "", "", "Coaches", "Display names Josh Dolan, Tim Slocombe, Joe Snyderwine. Sheet names and 'HC - Name' labels were left as they are."])
    if skipped:
        rows.append(["", "", "", "", "Names", "Formula cells were not changed: " + ", ".join(skipped)])

    out_dir = os.path.join(ROOT, "source", "corrected")
    dest = os.path.join(out_dir, f"{today} {TRACKER_NAME} (corrected).xlsx")
    print(f"{len(items)} cell edits; applying with Excel...")
    apply_with_excel(src, dest, [{**i, "new": i["new"]} for i in items], rows)
    unexpected, wrong, m1, m2, sheets = verify(src, dest, items)
    print(f"Saved {os.path.relpath(dest, ROOT)}")
    print(f"Pictures: {m1} in the original, {m2} in the copy. Sheets: {len(sheets)} (last: {sheets[-1]!r})")
    print(f"Listed edits landed: {len(items) - len(wrong)}/{len(items)}; other cells changed: {len(unexpected)}")
    for w in wrong[:20]:
        print("  WRONG", w)
    for u in unexpected[:30]:
        print("  CHANGED", u)

    with open(os.path.join(ROOT, "data", "corrected-tracker-changes.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write(f"# Corrected tracker: change list\n\nMade {today} from `{os.path.basename(src)}`. {len(items)} cell edits.\n"
                "The corrected copy itself stays on the owner's computer (source/corrected/).\n\n"
                "| Sheet | Cell | Old value | New value | Fix | Why |\n|---|---|---|---|---|---|\n")
        for i in items:
            esc = lambda v: str(shown(v)).replace("|", "\\|")
            f.write(f"| {i['sheet']} | {i['cell']} | {esc(i['old'])} | {esc(i['new'])} | {i['fix']} | {i['why']} |\n")
    sys.exit(1 if (wrong or unexpected) else 0)


if __name__ == "__main__":
    main()
