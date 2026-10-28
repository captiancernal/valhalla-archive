"""End-to-end spot check: read the tracker directly (read-only) and compare it with the BUILT website.

This deliberately does not reuse build_data.py: it re-reads raw cells its own way and looks for the
values in the rendered HTML in _site/, so a bug anywhere in the pipeline shows up as a mismatch.
Differences explained by an owner-approved fix are listed as "explained", not as failures.

  py -P scripts/spot_check.py [path/to/tracker.xlsx]     (default: newest source/*.xlsx)

Checks, for all three programs and every season in the tracker:
  - season record (W-L) and Big 12 finish, against the season page
  - the six season leaders on the coach sheet (name + number), against the season page
  - career record, against the program page
  - national champion and runner-up (Standings sheet), against the League Almanac year page
Writes data/spot-check.md; exits 1 if anything unexplained differs.
"""
import glob
import html
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SITE = os.path.join(ROOT, "_site")
SHEETS = {"CLEV": "Coach Josh Dolan", "OHIO": "Coach Tim Slocombe ", "M-OH": "Coach Joe Snyderwine"}
SLUG = {"CLEV": "cleveland-state", "OHIO": "ohio", "M-OH": "miami-oh"}
LEADERS = ["Passing Leader", "Rushing Leader", "Receiving Leader", "Tackle Leader", "Sack Leader", "INT Leader"]
# Owner-approved fixes that intentionally change what the tracker says (data/fixes.json)
EXPLAINED = {("CLEV", 2029, "rank"): "F26", ("OHIO", 2029, "rank"): "F26"}


def num(v):
    if v is None or v == "":
        return None
    try:
        f = float(str(v).replace(",", "").strip())
    except ValueError:
        return None
    return int(f) if f == int(f) else f


def fmt(n):
    if isinstance(n, float) and n != int(n):
        return f"{n:,.1f}"
    return f"{int(n):,}"


def clean_name(s):
    s = re.sub(r"\s+-\s+.*$", "", str(s))  # "Name - SR (RS)" or "Name - ROLB"
    s = re.sub(r"\s*\b(FR|SO|JR|SR)\b.*$", "", s).strip()  # "Name SO (RS)"
    return re.sub(r"\s+", " ", s)


def page_text(rel):
    p = os.path.join(SITE, rel, "index.html")
    if not os.path.exists(p):
        return None
    raw = open(p, encoding="utf-8").read()
    return raw, html.unescape(re.sub(r"<[^>]+>", " ", raw))


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(ROOT, "source", "*.xlsx")))[-1]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    aliases = json.load(open(os.path.join(ROOT, "data", "aliases.json"), encoding="utf-8"))
    canon = {v.lower(): k for k, vs in aliases["players"].items() for v in vs}
    teams = json.load(open(os.path.join(ROOT, "data", "reference", "teams.json"), encoding="utf-8"))["teams"]
    ok, bad, explained = [], [], []

    def record(good, what, detail="", key=None):
        if good:
            ok.append(what)
        elif key in EXPLAINED:
            explained.append(f"{what}: {detail} (approved fix {EXPLAINED[key]})")
        else:
            bad.append(f"{what}: {detail}")

    for pid, sheet in SHEETS.items():
        rows = list(wb[sheet].iter_rows(min_row=1, max_row=40, max_col=40, values_only=True))
        hdr = rows[1]
        col = {str(h).strip(): i for i, h in enumerate(hdr) if h}
        wi, li = col.get("W", col.get("WIN")), col.get("L", col.get("LOSS"))
        tw = tl = 0
        for r in rows[2:]:
            yr = num(r[1])
            w, l = num(r[wi]), num(r[li])
            if not yr or w is None:
                continue
            tw, tl = tw + w, tl + l
            rank = next((str(c).strip().lower() for c in r if c and re.fullmatch(r"\d{1,2}(st|nd|rd|th)", str(c).strip(), re.I)), None)
            pg = page_text(f"programs/{SLUG[pid]}/{yr}")
            if not pg:
                bad.append(f"{pid} {yr}: season page missing")
                continue
            raw, text = pg
            desc = re.search(r'<meta name="description" content="([^"]*)"', raw).group(1)
            record(f"{w}-{l}" in desc, f"{pid} {yr} record {w}-{l}", f"page says: {desc}")
            if rank:
                record(f"{rank} in the big 12" in desc.lower(), f"{pid} {yr} Big 12 finish {rank}", f"page says: {desc}", (pid, yr, "rank"))
            grid = raw[raw.find('class="leader-grid"'):]
            grid = html.unescape(re.sub(r"<[^>]+>", " ", grid[:grid.find("</div>\n\n") if "</div>\n\n" in grid else 6000]))
            for lab in LEADERS:
                if lab not in col:
                    continue
                name, val = r[col[lab]], num(r[col[lab] + 1])
                if not name or val is None:
                    continue
                n = clean_name(name)
                n = canon.get(n.lower(), n)
                record(n in grid and fmt(val) in grid, f"{pid} {yr} {lab.lower()} {n} {fmt(val)}", "not found on the season page")
        pg = page_text(f"programs/{SLUG[pid]}")
        record(pg and f"{tw}-{tl}" in pg[1], f"{pid} career record {tw}-{tl}", "not on the program page")

    st = list(wb["Standings"].iter_rows(min_row=2, max_row=4, max_col=30, values_only=True))
    for i, yr in enumerate(st[0]):
        yr = num(yr)
        if not yr or not st[1][i]:
            continue
        pg = page_text(f"almanac/{yr}")
        if not pg:
            bad.append(f"{yr}: almanac page missing")
            continue
        lead = pg[1][pg[1].find(f"The {yr} season"):][:400]
        champ, ru = teams.get(st[1][i], st[1][i]), teams.get(st[2][i], st[2][i])
        record(champ in lead and ru in lead and lead.find(champ) < lead.find(ru),
               f"{yr} national champion {champ} over {ru}", f"almanac says: {' '.join(lead.split())[:120]}")

    out = [f"# Spot check: tracker vs built site\n", f"Tracker: `{os.path.basename(src)}`\n",
           f"{len(ok)} matched, {len(explained)} explained by approved fixes, {len(bad)} unexplained.\n"]
    if bad:
        out += ["\n## Unexplained differences\n"] + [f"- {b}\n" for b in bad]
    if explained:
        out += ["\n## Explained by approved fixes\n"] + [f"- {e}\n" for e in explained]
    with open(os.path.join(ROOT, "data", "spot-check.md"), "w", encoding="utf-8", newline="\n") as f:
        f.writelines(out)
    print(f"{len(ok)} matched, {len(explained)} explained by approved fixes, {len(bad)} unexplained")
    for b in bad[:40]:
        print("  DIFF", b)
    for e in explained:
        print("  OK*", e)
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
