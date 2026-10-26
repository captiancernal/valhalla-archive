"""Dump non-empty cells of workbook sheets as text, one line per row (for inspection).

Read-only. Usage: py -P scripts/dump_sheet.py "<file>.xlsx" <out_dir> [sheet name ...]
"""
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def fmt(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\n", "\\n")
    return s


def main(path, out_dir, *names):
    wb = openpyxl.load_workbook(path, data_only=True)
    os.makedirs(out_dir, exist_ok=True)
    for ws in wb.worksheets:
        if names and ws.title not in names:
            continue
        lines = []
        for row in ws.iter_rows():
            cells = [f"{get_column_letter(c.column)}={fmt(c.value)}" for c in row if c.value is not None]
            if cells:
                lines.append(f"R{row[0].row}: " + " | ".join(cells))
        if ws.merged_cells.ranges:
            lines.append("MERGED: " + " ".join(sorted(str(r) for r in ws.merged_cells.ranges)[:400]))
        if ws.tables:
            lines.append("TABLES: " + " ".join(f"{n}={ws.tables[n].ref}" for n in ws.tables))
        fn = re.sub(r"[^A-Za-z0-9]+", "_", ws.title).strip("_") + ".txt"
        with open(os.path.join(out_dir, fn), "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        print(f"{ws.title}: {len(lines)} lines -> {fn}")


if __name__ == "__main__":
    main(*sys.argv[1:])
