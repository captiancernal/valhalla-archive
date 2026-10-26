"""Inventory the tracker workbook: sheets, sizes, Excel tables, images, merged cells.

Read-only. Usage: py -P scripts/inventory.py "source/<file>.xlsx" [out.json]
"""
import json
import sys
from collections import Counter

import openpyxl


def main(path, out=None):
    wb = openpyxl.load_workbook(path, data_only=True)  # cached values, never formulas
    report = []
    for ws in wb.worksheets:
        kinds = Counter()
        errors = Counter()
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                kinds[type(v).__name__] += 1
                if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")):
                    errors[v] += 1
        report.append({
            "sheet": ws.title,
            "dims": ws.dimensions,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "tables": {name: ws.tables[name].ref for name in ws.tables},
            "images": len(getattr(ws, "_images", [])),
            "merged": len(ws.merged_cells.ranges),
            "value_types": dict(kinds),
            "error_cells": dict(errors),
        })
    for r in report:
        print(f"{r['sheet']:<28} {r['dims']:<14} tables={len(r['tables']):<3} imgs={r['images']:<4} "
              f"merged={r['merged']:<4} types={r['value_types']} errs={r['error_cells']}")
    if out:
        with open(out, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2, default=str)


if __name__ == "__main__":
    main(*sys.argv[1:])
